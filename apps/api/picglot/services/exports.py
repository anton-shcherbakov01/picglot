"""Export generation for every supported output format."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime, timedelta
from typing import Any

from PIL import Image
from sqlalchemy.orm import Session

from picglot.core.config import settings
from picglot.core.errors import AppError, ErrorCode
from picglot.core.logging import get_logger
from picglot.db.models import DocumentPage, Export, Project, TextRegion
from picglot.domain.enums import AssetKind, ExportFormat, RegionType
from picglot.services import projects as project_service
from picglot.vision import pdf as pdf_tools
from picglot.vision import render as render_tools
from picglot.vision.preprocess import encode

log = get_logger(__name__)

SCHEMA_VERSION = "1.0"


@dataclass(slots=True)
class ExportOptions:
    #: "source" | "translation" | "bilingual"
    content: str = "translation"
    quality: int = 92
    scale: float = 1.0
    keep_line_breaks: bool = True
    include_confidence: bool = False
    include_coordinates: bool = False
    watermark: str | None = None
    password: str | None = None
    pdfa: bool = False
    pdf_layout: str = "sequential"
    sheet_per_page: bool = True
    include_source_image_sheet: bool = False
    strip_metadata: bool = True

    @classmethod
    def from_dict(cls, data: dict[str, Any] | None) -> ExportOptions:
        options = cls()
        for key, value in (data or {}).items():
            if hasattr(options, key) and value is not None:
                setattr(options, key, value)
        return options

    def as_dict(self) -> dict[str, Any]:
        return {
            key: getattr(self, key)
            for key in (
                "content",
                "quality",
                "scale",
                "keep_line_breaks",
                "include_confidence",
                "include_coordinates",
                "watermark",
                "pdfa",
                "pdf_layout",
                "sheet_per_page",
                "include_source_image_sheet",
                "strip_metadata",
            )
        }


def build(
    session: Session,
    project: Project,
    fmt: ExportFormat | str,
    options: ExportOptions | None = None,
) -> Export:
    """Produce (or reuse) an export. Re-exporting identical settings is free."""
    fmt = ExportFormat(str(fmt))
    options = options or ExportOptions()

    from picglot.services.pipeline import settings_hash

    fingerprint = settings_hash(
        {**options.as_dict(), "version": project.document_version, "format": str(fmt)}
    )
    cached = project_service.find_cached_export(session, project.id, str(fmt), fingerprint)
    if cached is not None:
        log.info("export.cache_hit", project_id=project.id, format=str(fmt))
        return cached

    pages = sorted(project.pages, key=lambda page: page.page_number)
    if not pages:
        raise AppError(code=ErrorCode.EXPORT_FAILED, internal="project has no pages")

    builder = _BUILDERS.get(fmt)
    if builder is None:
        raise AppError(code=ErrorCode.EXPORT_FAILED, details={"format": str(fmt)})

    try:
        data, extension = builder(project, pages, options)
    except AppError:
        raise
    except Exception as exc:  # pragma: no cover - defensive
        log.exception("export.failed", project_id=project.id, format=str(fmt))
        raise AppError(
            code=ErrorCode.EXPORT_FAILED,
            details={"format": str(fmt)},
            internal=f"{type(exc).__name__}: {exc}",
        ) from exc

    asset = project_service.store_asset(
        session,
        project,
        data=data,
        kind=AssetKind.EXPORT,
        mime_type=fmt.media_type,
        extension=extension,
        original_filename=f"{_safe_stem(project.name)}.{extension}",
        metadata={"format": str(fmt), "settings": options.as_dict()},
    )
    export = Export(
        project_id=project.id,
        format=str(fmt),
        asset_id=asset.id,
        settings=options.as_dict(),
        settings_hash=fingerprint,
        byte_size=len(data),
        expires_at=project.expires_at or datetime.now(UTC) + timedelta(days=7),
    )
    session.add(export)
    session.flush()
    log.info("export.created", project_id=project.id, format=str(fmt), bytes=len(data))
    return export


# --------------------------------------------------------------------------- #
# Text-shaped exports
# --------------------------------------------------------------------------- #
def _text_of(region: TextRegion, options: ExportOptions, target_language: str | None) -> str:
    if options.content == "source":
        return region.normalized_text or region.source_text or ""
    translation = next(
        (
            item
            for item in region.translations
            if item.is_active and (not target_language or item.target_language == target_language)
        ),
        None,
    )
    translated = translation.translated_text if translation else ""
    source = region.normalized_text or region.source_text or ""
    if options.content == "bilingual":
        return f"{source}\n{translated}" if translated else source
    return translated or source


def _ordered_regions(page: DocumentPage) -> list[TextRegion]:
    return sorted(page.regions, key=lambda region: (region.reading_order, region.id))


def _build_txt(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    blocks: list[str] = []
    for page in pages:
        if len(pages) > 1:
            blocks.append(f"--- page {page.page_number} ---")
        for region in _ordered_regions(page):
            text = _text_of(region, options, project.target_language).strip()
            if text:
                blocks.append(text)
        blocks.append("")
    separator = "\n" if options.keep_line_breaks else " "
    return separator.join(blocks).strip().encode("utf-8"), "txt"


def _build_markdown(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    lines: list[str] = [f"# {project.name}", ""]
    for page in pages:
        if len(pages) > 1:
            lines.append(f"## Page {page.page_number}")
            lines.append("")
        for region in _ordered_regions(page):
            text = _text_of(region, options, project.target_language).strip()
            if not text:
                continue
            region_type = RegionType(region.region_type)
            if region_type is RegionType.HEADING:
                lines.append(f"### {text}")
            elif region_type is RegionType.LIST_ITEM:
                lines.append(f"- {text.lstrip('•·-* ')}")
            elif region_type is RegionType.FOOTNOTE:
                lines.append(f"> {text}")
            else:
                lines.append(text)
            lines.append("")
        for table in page.tables:
            lines.extend(_markdown_table(table))
            lines.append("")
    return "\n".join(lines).encode("utf-8"), "md"


def _markdown_table(table: Any) -> list[str]:
    grid: dict[tuple[int, int], str] = {
        (cell.row, cell.col): (cell.raw_text or "").replace("|", "\\|") for cell in table.cells
    }
    rows: list[str] = []
    for row in range(table.row_count):
        cells = [grid.get((row, col), "") for col in range(table.column_count)]
        rows.append("| " + " | ".join(cells) + " |")
        if row == 0 and table.has_header:
            rows.append("| " + " | ".join(["---"] * table.column_count) + " |")
    return rows


def _build_json(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    payload: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "project": {
            "id": project.id,
            "name": project.name,
            "tool": project.tool_type,
            "source_language": project.source_language,
            "target_language": project.target_language,
            "created_at": project.created_at.isoformat(),
            "quality": {
                "score": project.quality_score,
                "band": project.quality_band,
                "reasons": project.quality_reasons,
            },
        },
        "pages": [],
    }
    for page in pages:
        page_payload: dict[str, Any] = {
            "number": page.page_number,
            "width": page.width,
            "height": page.height,
            "rotation": page.rotation,
            "detected_language": page.detected_language,
            "confidence": page.ocr_confidence,
            "regions": [],
            "tables": [],
        }
        for region in _ordered_regions(page):
            translation = next((item for item in region.translations if item.is_active), None)
            entry: dict[str, Any] = {
                "id": region.id,
                "type": region.region_type,
                "reading_order": region.reading_order,
                "source_text": region.normalized_text or region.source_text,
                "raw_ocr_text": region.source_text,
                "translated_text": translation.translated_text if translation else None,
                "detected_language": region.detected_language,
            }
            if options.include_confidence:
                entry["confidence"] = region.confidence
                entry["low_confidence_spans"] = region.low_confidence_spans
            if options.include_coordinates:
                entry["bounding_box"] = region.bounding_box
                entry["polygon"] = region.polygon
                entry["rotation"] = region.rotation
                entry["style"] = region.style
            page_payload["regions"].append(entry)

        for table in page.tables:
            page_payload["tables"].append(
                {
                    "index": table.index_on_page,
                    "rows": table.row_count,
                    "cols": table.column_count,
                    "has_header": table.has_header,
                    "structure_ambiguous": table.structure_ambiguous,
                    "confidence": table.confidence,
                    "cells": [
                        {
                            "row": cell.row,
                            "col": cell.col,
                            "row_span": cell.row_span,
                            "col_span": cell.col_span,
                            "text": cell.raw_text,
                            "type": cell.value_type,
                            "numeric_value": cell.numeric_value,
                            "currency": cell.currency,
                        }
                        for cell in table.cells
                    ],
                }
            )
        payload["pages"].append(page_payload)

    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8"), "json"


def _build_csv(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    buffer = io.StringIO(newline="")
    writer = csv.writer(buffer, quoting=csv.QUOTE_MINIMAL)

    all_tables = [table for page in pages for table in page.tables]
    if all_tables:
        for index, table in enumerate(all_tables):
            if index:
                writer.writerow([])
            grid: dict[tuple[int, int], str] = {
                (cell.row, cell.col): cell.raw_text or "" for cell in table.cells
            }
            for row in range(table.row_count):
                writer.writerow([grid.get((row, col), "") for col in range(table.column_count)])
    else:
        header = ["page", "reading_order", "type", "source_text", "translated_text"]
        if options.include_confidence:
            header.append("confidence")
        if options.include_coordinates:
            header += ["x", "y", "width", "height"]
        writer.writerow(header)
        for page in pages:
            for region in _ordered_regions(page):
                translation = next((item for item in region.translations if item.is_active), None)
                row: list[Any] = [
                    page.page_number,
                    region.reading_order,
                    region.region_type,
                    region.normalized_text or region.source_text,
                    translation.translated_text if translation else "",
                ]
                if options.include_confidence:
                    row.append(region.confidence)
                if options.include_coordinates:
                    box = region.bounding_box or {}
                    row += [box.get("x"), box.get("y"), box.get("width"), box.get("height")]
                writer.writerow(row)

    # BOM so Excel opens UTF-8 correctly on Windows.
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8"), "csv"


# --------------------------------------------------------------------------- #
# Office formats
# --------------------------------------------------------------------------- #
def _build_docx(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()
    document.core_properties.title = project.name
    document.core_properties.author = settings.brand_name

    alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }

    for page_index, page in enumerate(pages):
        if page_index:
            document.add_page_break()

        for region in _ordered_regions(page):
            text = _text_of(region, options, project.target_language).strip()
            if not text:
                continue
            region_type = RegionType(region.region_type)
            style = region.style or {}

            if region_type is RegionType.HEADING:
                paragraph = document.add_heading(text, level=2)
            elif region_type is RegionType.LIST_ITEM:
                paragraph = document.add_paragraph(text.lstrip("•·-* "), style="List Bullet")
            else:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(text)
                run.bold = bool(style.get("bold"))
                run.italic = bool(style.get("italic"))
                run.underline = bool(style.get("underline"))
                size = style.get("font_size")
                if isinstance(size, (int, float)) and size > 0:
                    run.font.size = Pt(max(6, min(72, float(size) * 0.75)))
                colour = str(style.get("color") or "").lstrip("#")
                if len(colour) == 6:
                    try:
                        run.font.color.rgb = RGBColor.from_string(colour.upper())
                    except ValueError:
                        pass
            align = alignment.get(str(style.get("align", "left")))
            if align is not None:
                paragraph.alignment = align

        for table_record in page.tables:
            if table_record.row_count < 1 or table_record.column_count < 1:
                continue
            table = document.add_table(rows=table_record.row_count, cols=table_record.column_count)
            table.style = "Table Grid"
            for cell in table_record.cells:
                if cell.row < table_record.row_count and cell.col < table_record.column_count:
                    target = table.cell(cell.row, cell.col)
                    target.text = cell.raw_text or ""
                    if cell.is_header:
                        for paragraph in target.paragraphs:
                            for run in paragraph.runs:
                                run.bold = True
            document.add_paragraph()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), "docx"


def _build_xlsx(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)

    tables = [(page, table) for page in pages for table in page.tables]
    if not tables:
        # No table detected: fall back to a readable text dump per page.
        for page in pages:
            sheet = workbook.create_sheet(f"Page {page.page_number}"[:31])
            sheet.append(["Reading order", "Type", "Text", "Confidence"])
            for cell in sheet[1]:
                cell.font = header_font
            for region in _ordered_regions(page):
                sheet.append(
                    [
                        region.reading_order,
                        region.region_type,
                        _text_of(region, options, project.target_language),
                        round(region.confidence, 3) if region.confidence else None,
                    ]
                )
            sheet.column_dimensions["C"].width = 80
            sheet.column_dimensions["B"].width = 16
    else:
        for index, (page, table) in enumerate(tables):
            title = (
                f"P{page.page_number}-T{table.index_on_page + 1}"
                if options.sheet_per_page
                else "Data"
            )
            sheet = (
                workbook.create_sheet(title[:31])
                if options.sheet_per_page or index == 0
                else workbook[workbook.sheetnames[0]]
            )
            # Fixed before the loop: `max_row` grows as cells are written, so
            # reading it per cell would shift every subsequent row.
            # openpyxl reports max_row == 1 for an empty sheet.
            base_row = 1 if sheet.max_row == 1 else sheet.max_row + 2

            widths: dict[int, int] = {}
            for cell in table.cells:
                target = sheet.cell(row=base_row + cell.row, column=cell.col + 1)
                value: Any = cell.raw_text or ""
                if cell.value_type in {"number", "currency", "percent"} and (
                    cell.numeric_value is not None
                ):
                    value = cell.numeric_value
                    if cell.value_type == "currency" and cell.currency:
                        target.number_format = _currency_format(cell.currency)
                    elif cell.value_type == "percent":
                        target.number_format = "0.0%"
                elif cell.value_type == "formula" and (cell.raw_text or "").startswith("="):
                    value = cell.raw_text
                target.value = value
                target.border = border
                target.alignment = Alignment(
                    vertical="center",
                    horizontal="right" if cell.value_type != "text" else "left",
                    wrap_text=True,
                )
                if cell.is_header:
                    target.font = header_font
                if options.include_confidence and cell.confidence is not None:
                    target.comment = Comment(
                        f"OCR confidence: {cell.confidence:.0%}", settings.brand_name
                    )
                widths[cell.col + 1] = max(
                    widths.get(cell.col + 1, 10), min(50, len(str(value)) + 4)
                )
                if cell.row_span > 1 or cell.col_span > 1:
                    try:
                        sheet.merge_cells(
                            start_row=target.row,
                            start_column=target.column,
                            end_row=target.row + cell.row_span - 1,
                            end_column=target.column + cell.col_span - 1,
                        )
                    except ValueError:
                        pass
            for column, width in widths.items():
                sheet.column_dimensions[get_column_letter(column)].width = width

            if table.structure_ambiguous:
                note = sheet.cell(row=1, column=table.column_count + 2)
                note.value = "Structure uncertain — please verify rows and columns"
                note.font = Font(italic=True, color="B45309")

    if options.include_source_image_sheet:
        _attach_source_images(workbook, pages)

    if not workbook.sheetnames:
        workbook.create_sheet("Empty")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), "xlsx"


def _currency_format(code: str) -> str:
    symbol = {"USD": "$", "EUR": "€", "GBP": "£", "RUB": "₽", "JPY": "¥"}.get(code, "")
    return f"{symbol}#,##0.00" if symbol else "#,##0.00"


def _attach_source_images(workbook: Any, pages: list[DocumentPage]) -> None:
    from openpyxl.drawing.image import Image as XlsxImage

    for page in pages[:5]:
        asset = _asset_by_id(page, page.preview_asset_id)
        image = project_service.load_image_asset(asset)
        if image is None:
            continue
        sheet = workbook.create_sheet(f"Source {page.page_number}"[:31])
        buffer = io.BytesIO()
        image.convert("RGB").save(buffer, "PNG")
        buffer.seek(0)
        sheet.add_image(XlsxImage(buffer), "A1")


# --------------------------------------------------------------------------- #
# Image and PDF
# --------------------------------------------------------------------------- #
def _page_image(page: DocumentPage, options: ExportOptions) -> Image.Image | None:
    preference = (
        [page.rendered_asset_id, page.normalized_asset_id]
        if options.content != "source"
        else [page.normalized_asset_id, page.rendered_asset_id]
    )
    for asset_id in preference:
        image = project_service.load_image_asset(_asset_by_id(page, asset_id))
        if image is not None:
            if options.scale and options.scale != 1.0:
                image = image.resize(
                    (
                        max(1, int(image.width * options.scale)),
                        max(1, int(image.height * options.scale)),
                    ),
                    Image.Resampling.LANCZOS,
                )
            if options.watermark:
                image = render_tools.add_watermark(image, options.watermark)
            return image
    return None


def _asset_by_id(page: DocumentPage, asset_id: str | None) -> Any:
    if not asset_id:
        return None
    for asset in page.project.assets:
        if asset.id == asset_id:
            return asset
    return None


def _build_image(fmt: ExportFormat):
    def builder(
        project: Project, pages: list[DocumentPage], options: ExportOptions
    ) -> tuple[bytes, str]:
        images = [image for page in pages if (image := _page_image(page, options))]
        if not images:
            raise AppError(code=ErrorCode.EXPORT_FAILED, internal="no renderable page image")
        if len(images) == 1:
            pillow_format = {"png": "PNG", "jpg": "JPEG", "webp": "WEBP"}[fmt.extension]
            return encode(images[0], pillow_format, options.quality), fmt.extension
        # Several pages in a single-image format become a ZIP of images.
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            pillow_format = {"png": "PNG", "jpg": "JPEG", "webp": "WEBP"}[fmt.extension]
            for page, image in zip(pages, images, strict=False):
                archive.writestr(
                    f"page-{page.page_number:03d}.{fmt.extension}",
                    encode(image, pillow_format, options.quality),
                )
        return buffer.getvalue(), "zip"

    return builder


def _build_pdf(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    images = [image for page in pages if (image := _page_image(page, options))]
    if not images:
        raise AppError(code=ErrorCode.EXPORT_FAILED, internal="no page images")
    data = pdf_tools.images_to_pdf(
        images,
        quality=options.quality,
        metadata={"title": project.name, "producer": settings.brand_name},
    )
    return _finish_pdf(data, options), "pdf"


def _build_searchable_pdf(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    payload: list[tuple[Image.Image, list[Any]]] = []
    for page in pages:
        image = _page_image(page, options)
        if image is None:
            continue
        regions = [project_service.region_to_domain(row) for row in _ordered_regions(page)]
        if options.content != "source":
            for region, row in zip(regions, _ordered_regions(page), strict=False):
                translation = next((item for item in row.translations if item.is_active), None)
                if translation and translation.translated_text:
                    region.normalized_text = translation.translated_text
        payload.append((image, regions))

    if not payload:
        raise AppError(code=ErrorCode.EXPORT_FAILED, internal="no pages for searchable PDF")

    data = pdf_tools.searchable_pdf(
        payload,
        quality=options.quality,
        metadata={"title": project.name},
        language=project.target_language or project.source_language,
    )
    return _finish_pdf(data, options), "pdf"


def _build_bilingual_pdf(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    pairs: list[tuple[Image.Image, Image.Image]] = []
    for page in pages:
        original = project_service.load_image_asset(_asset_by_id(page, page.normalized_asset_id))
        translated = project_service.load_image_asset(_asset_by_id(page, page.rendered_asset_id))
        if original is None:
            continue
        pairs.append((original, translated or original))
    if not pairs:
        raise AppError(code=ErrorCode.EXPORT_FAILED, internal="no page pairs")
    data = pdf_tools.bilingual_pdf(pairs, layout=options.pdf_layout, quality=options.quality)
    return _finish_pdf(data, options), "pdf"


def _finish_pdf(data: bytes, options: ExportOptions) -> bytes:
    if options.pdfa:
        data, converted = pdf_tools.to_pdfa(data)
        if not converted:
            log.info("export.pdfa_unavailable")
    if options.password:
        data = pdf_tools.set_password(data, options.password)
    return data


def _build_zip(
    project: Project, pages: list[DocumentPage], options: ExportOptions
) -> tuple[bytes, str]:
    """Everything the project can produce, in one archive."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for page in pages:
            image = _page_image(page, options)
            if image is not None:
                archive.writestr(f"images/page-{page.page_number:03d}.png", encode(image, "PNG"))
        text_data, _ = _build_txt(project, pages, options)
        archive.writestr("text.txt", text_data)
        json_data, _ = _build_json(project, pages, options)
        archive.writestr("data.json", json_data)
        if any(page.tables for page in pages):
            csv_data, _ = _build_csv(project, pages, options)
            archive.writestr("tables.csv", csv_data)
        archive.writestr(
            "README.txt",
            (
                f"{settings.brand_name} export\n"
                f"Project: {project.name}\n"
                f"Created: {datetime.now(UTC).isoformat()}\n"
                f"Pages: {len(pages)}\n"
                f"Source language: {project.source_language or 'auto'}\n"
                f"Target language: {project.target_language or '-'}\n"
            ).encode(),
        )
    return buffer.getvalue(), "zip"


_BUILDERS = {
    ExportFormat.TXT: _build_txt,
    ExportFormat.MARKDOWN: _build_markdown,
    ExportFormat.JSON: _build_json,
    ExportFormat.CSV: _build_csv,
    ExportFormat.DOCX: _build_docx,
    ExportFormat.XLSX: _build_xlsx,
    ExportFormat.PNG: _build_image(ExportFormat.PNG),
    ExportFormat.JPG: _build_image(ExportFormat.JPG),
    ExportFormat.WEBP: _build_image(ExportFormat.WEBP),
    ExportFormat.PDF: _build_pdf,
    ExportFormat.PDF_SEARCHABLE: _build_searchable_pdf,
    ExportFormat.PDF_BILINGUAL: _build_bilingual_pdf,
    ExportFormat.ZIP: _build_zip,
}


def _safe_stem(name: str) -> str:
    from picglot.services.files import safe_filename

    stem = safe_filename(name, fallback="export")
    return stem.rsplit(".", 1)[0][:80] or "export"


def download_url(session: Session, export: Export, *, filename: str | None = None) -> str:
    from picglot.services import storage

    asset = project_service.get_asset(session, export.asset_id) if export.asset_id else None
    if asset is None:
        raise AppError(code=ErrorCode.NOT_FOUND, internal="export asset missing")
    export.download_count = int(export.download_count) + 1
    return storage.get_storage().signed_download_url(
        asset.storage_key,
        filename=filename or asset.original_filename,
        content_type=asset.mime_type,
    )


def batch_report_csv(rows: list[dict[str, Any]]) -> bytes:
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(
        buffer,
        fieldnames=["file", "status", "pages", "error_code", "project_id", "download"],
        extrasaction="ignore",
    )
    writer.writeheader()
    for row in rows:
        writer.writerow(row)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")
